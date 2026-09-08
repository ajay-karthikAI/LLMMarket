"""
LLMarket — Healthcare AI Job Intelligence Dashboard
"""

import logging
import sys
from datetime import datetime
from pathlib import Path

import base64

import pandas as pd
import plotly.express as px
import streamlit as st
import streamlit.components.v1 as components

ROOT = Path(__file__).parent
sys.path.insert(0, str(ROOT))

from utils.logging_config import setup_logging
from database.db import init_db, get_all_jobs
from extraction.location import get_canonical_locations
from analytics.skills_analysis import (
    skill_frequency, top_paying_skills, skill_combinations,
    study_roadmap, skills_summary_stats,
)
from analytics.trends import skill_trend_df, fastest_growing_skills
from jobs.fetcher import COMPANY_REGISTRY, run_all_fetchers

setup_logging()
logger = logging.getLogger(__name__)
init_db()

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="LLMarket",
    page_icon="⬡",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── CSS ───────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&family=JetBrains+Mono:wght@400;500&display=swap');

*, *::before, *::after { box-sizing: border-box; }
html, body, [class*="css"] { font-family: 'Inter', sans-serif !important; }

/* ── Chrome removal ─────────────────────────────────────────────────────── */
#MainMenu, footer, header { visibility: hidden; }
.stDeployButton,
[data-testid="stToolbar"],
[data-testid="stSidebarCollapseButton"] { display: none !important; }

/* ── Backgrounds ─────────────────────────────────────────────────────────── */
.stApp, [data-testid="stAppViewContainer"] { background: #0b0d0c; }
.main .block-container {
    background: #0b0d0c;
    padding: 1.5rem 2.25rem 3rem !important;
    max-width: 100% !important;
}

/* ── Sidebar shell ───────────────────────────────────────────────────────── */
[data-testid="stSidebar"] {
    background: #0d0f0e !important;
    border-right: 1px solid #192020 !important;
    min-width: 210px !important;
    max-width: 210px !important;
}
[data-testid="stSidebar"] > div,
section[data-testid="stSidebar"] .block-container { padding: 0 !important; }

/* ── Nav buttons ─────────────────────────────────────────────────────────── */
.nav-wrap button {
    background: transparent !important;
    border: none !important;
    border-left: 3px solid transparent !important;
    border-radius: 0 !important;
    color: #4a6454 !important;
    font-size: 0.845rem !important;
    font-weight: 500 !important;
    text-align: left !important;
    padding: 0.6rem 1.25rem !important;
    width: 100% !important;
    transition: color 0.12s, background 0.12s !important;
    box-shadow: none !important;
}
.nav-wrap button:hover {
    color: #a0c8b0 !important;
    background: #111714 !important;
    border-left-color: transparent !important;
}
.nav-active button {
    background: #101f15 !important;
    border-left: 3px solid #3ddc84 !important;
    color: #3ddc84 !important;
    font-weight: 600 !important;
}
.nav-active button:hover { color: #3ddc84 !important; }

/* ── Refresh button ──────────────────────────────────────────────────────── */
.refresh-btn button {
    background: #101f15 !important;
    border: 1px solid #1e3326 !important;
    color: #3ddc84 !important;
    border-radius: 8px !important;
    font-size: 0.82rem !important;
    font-weight: 600 !important;
    width: 100% !important;
    box-shadow: none !important;
}
.refresh-btn button:hover {
    background: #162a1e !important;
    border-color: #3ddc84 !important;
}

/* ── Metrics ─────────────────────────────────────────────────────────────── */
[data-testid="stMetric"] {
    background: #0f1310;
    border: 1px solid #181f18;
    border-radius: 8px;
    padding: 1.25rem 1.5rem !important;
}
[data-testid="stMetricLabel"] {
    font-size: 0.67rem !important;
    font-weight: 700 !important;
    letter-spacing: 0.12em !important;
    text-transform: uppercase !important;
    color: #34503c !important;
}
[data-testid="stMetricValue"] {
    font-size: 2.1rem !important;
    font-weight: 700 !important;
    color: #dce8e0 !important;
    font-family: 'JetBrains Mono', monospace !important;
    line-height: 1.1 !important;
}
[data-testid="stMetricDelta"] {
    font-size: 0.71rem !important;
    color: #344c3c !important;
}
.salary-metric [data-testid="stMetricValue"] { color: #3ddc84 !important; }
.salary-metric [data-testid="stMetricDelta"] { color: #2a5a3a !important; }

/* ── Inputs ──────────────────────────────────────────────────────────────── */
[data-testid="stTextInput"] input {
    background: #111714 !important;
    border: 1px solid #1e2822 !important;
    border-radius: 10px !important;
    color: #c8d8cc !important;
    font-size: 0.88rem !important;
}
[data-testid="stTextInput"] input:focus {
    border-color: #3ddc84 !important;
    box-shadow: 0 0 0 2px rgba(61,220,132,0.1) !important;
}

/* ── Multiselect ─────────────────────────────────────────────────────────── */
[data-testid="stMultiSelect"] > div > div {
    background: #111714 !important;
    border: 1px solid #1e2822 !important;
    border-radius: 8px !important;
    color: #c8d8cc !important;
    font-size: 0.82rem !important;
}

/* ── Slider ──────────────────────────────────────────────────────────────── */
[data-testid="stSlider"] { padding: 0 0.25rem !important; }

/* ── Selectbox ───────────────────────────────────────────────────────────── */
[data-testid="stSelectbox"] > div > div {
    background: #111714 !important;
    border: 1px solid #1e2822 !important;
    border-radius: 8px !important;
    color: #c8d8cc !important;
    font-size: 0.85rem !important;
}

/* ── Generic buttons ─────────────────────────────────────────────────────── */
.stButton > button {
    background: transparent !important;
    border: 1px solid #1e2a22 !important;
    color: #7aaa8a !important;
    border-radius: 8px !important;
    font-size: 0.82rem !important;
    transition: all 0.12s !important;
    box-shadow: none !important;
}
.stButton > button:hover {
    border-color: #3ddc84 !important;
    color: #3ddc84 !important;
    background: rgba(61,220,132,0.04) !important;
}

/* ── Tabs ────────────────────────────────────────────────────────────────── */
.stTabs [data-baseweb="tab-list"] {
    background: transparent !important;
    border-bottom: 1px solid #1a2418 !important;
    gap: 2px !important;
}
.stTabs [data-baseweb="tab"] {
    background: transparent !important;
    color: #4a6050 !important;
    font-size: 0.82rem !important;
    font-weight: 500 !important;
    padding: 0.5rem 1rem !important;
    border: none !important;
}
.stTabs [aria-selected="true"] {
    color: #3ddc84 !important;
    background: rgba(61,220,132,0.04) !important;
    border-bottom: 2px solid #3ddc84 !important;
}

/* ── Scrollbar ───────────────────────────────────────────────────────────── */
::-webkit-scrollbar { width: 4px; height: 4px; }
::-webkit-scrollbar-track { background: #0b0d0c; }
::-webkit-scrollbar-thumb { background: #1e2a22; border-radius: 4px; }
</style>
""", unsafe_allow_html=True)

# ── Component helpers ─────────────────────────────────────────────────────────

_PALETTE = [
    "#3ddc84","#00b4d8","#f72585","#7b2d8b",
    "#ff6b35","#ffd60a","#06d6a0","#118ab2","#ef476f","#8338ec",
]

def _logo(company: str, size: int = 42) -> str:
    letter = company[0].upper()
    color = _PALETTE[hash(company) % len(_PALETTE)]
    return (
        f'<div style="width:{size}px;height:{size}px;border-radius:9px;flex-shrink:0;'
        f'background:{color}18;border:1px solid {color}40;'
        f'display:flex;align-items:center;justify-content:center;'
        f'font-size:{int(size/2.3)}px;font-weight:700;color:{color};'
        f'font-family:\'JetBrains Mono\',monospace;">{letter}</div>'
    )

def _time_ago(posted_at) -> str:
    if not posted_at:
        return ""
    try:
        dt = datetime.fromisoformat(str(posted_at).replace("Z", ""))
        h = int((datetime.now() - dt).total_seconds() // 3600)
        if h < 1:   return "just now"
        if h < 24:  return f"{h}h ago"
        d = h // 24
        if d < 7:   return f"{d}d ago"
        return f"{d//7}w ago"
    except Exception:
        return ""

def _salary_str(mn, mx) -> str:
    if not mn and not mx: return ""
    if mn == mx or not mx: return f"${mn:,.0f}"
    return f"${mn:,.0f}–${mx:,.0f}"

def _skill_tag(s: str) -> str:
    return (
        f'<span style="display:inline-block;padding:2px 8px;background:#0f1f16;'
        f'border:1px solid #1e3326;border-radius:20px;font-size:0.71rem;'
        f'color:#5aaa78;margin:2px;font-family:\'JetBrains Mono\',monospace;">{s}</span>'
    )

def _to_excel(jobs: list[dict]) -> bytes:
    import io
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Parse each job's skill set
    job_skill_sets: list[set[str]] = []
    for j in jobs:
        raw = j.get("skills") or ""
        job_skill_sets.append({s.strip() for s in raw.split("|") if s.strip()})

    # Collect all unique skills, ordered by frequency (most common first)
    from collections import Counter
    skill_freq: Counter = Counter()
    for ss in job_skill_sets:
        skill_freq.update(ss)
    all_skills = [s for s, _ in skill_freq.most_common()]

    # Build rows: fixed info columns + one boolean column per skill
    fixed_cols = ["Company", "Title", "Seniority", "Category",
                  "Location", "Remote", "Salary Min", "Salary Max", "Posted", "URL"]
    rows = []
    for j, skill_set in zip(jobs, job_skill_sets):
        row = {
            "Company":    j.get("company_name", ""),
            "Title":      j.get("title", ""),
            "Seniority":  j.get("seniority", ""),
            "Category":   j.get("role_category", ""),
            "Location":   j.get("location_normalized") or j.get("location", ""),
            "Remote":     "Yes" if j.get("is_remote") else "No",
            "Salary Min": j.get("salary_min"),
            "Salary Max": j.get("salary_max"),
            "Posted":     (j.get("posted_at") or "")[:10],
            "URL":        j.get("url", ""),
        }
        for skill in all_skills:
            row[skill] = 1 if skill in skill_set else ""
        rows.append(row)

    df = pd.DataFrame(rows, columns=fixed_cols + all_skills)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Jobs")
        ws = writer.sheets["Jobs"]

        # Style header row
        header_fill  = PatternFill("solid", fgColor="0F1F16")
        header_font  = Font(bold=True, color="3DDC84", size=10)
        skill_fill   = PatternFill("solid", fgColor="0A1A10")
        skill_font   = Font(bold=True, color="5AAA78", size=9)
        center       = Alignment(horizontal="center", vertical="center", wrap_text=False)
        thin_border  = Border(bottom=Side(style="thin", color="1E3326"))

        n_fixed = len(fixed_cols)
        for col_idx, cell in enumerate(ws[1], start=1):
            is_skill = col_idx > n_fixed
            cell.fill       = skill_fill if is_skill else header_fill
            cell.font       = skill_font if is_skill else header_font
            cell.alignment  = center if is_skill else Alignment(vertical="center")
            cell.border     = thin_border

        # Mark skill cells with a green tint where value == 1
        check_fill = PatternFill("solid", fgColor="0F2E1A")
        check_font = Font(color="3DDC84", bold=True, size=10)
        for row in ws.iter_rows(min_row=2, min_col=n_fixed + 1):
            for cell in row:
                if cell.value == 1:
                    cell.value     = "✓"
                    cell.fill      = check_fill
                    cell.font      = check_font
                    cell.alignment = center

        # Column widths
        fixed_widths = [22, 42, 14, 20, 16, 7, 12, 12, 11, 48]
        for i, w in enumerate(fixed_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for i in range(n_fixed + 1, len(all_skills) + n_fixed + 1):
            ws.column_dimensions[get_column_letter(i)].width = 14

        # Freeze the fixed columns so skills scroll independently
        ws.freeze_panes = "K2"

        # Auto-filter on every column so clicking headers lets you filter
        ws.auto_filter.ref = ws.dimensions

    return buf.getvalue()


def _job_card(job: dict, show_skills: bool = False) -> str:
    company  = job.get("company_name", "?")
    title    = job.get("title", "Untitled")
    url      = job.get("url", "#")
    loc      = job.get("location_normalized") or job.get("location") or "—"
    cat      = job.get("role_category") or "AI / ML"
    seniority = job.get("seniority") or ""
    time_str = _time_ago(job.get("posted_at"))
    sal      = _salary_str(job.get("salary_min"), job.get("salary_max"))

    badge = (
        f'<span style="font-size:0.63rem;color:#486050;background:#0d1a12;'
        f'border:1px solid #1a2e1e;border-radius:4px;padding:1px 7px;margin-left:6px;">'
        f'{seniority}</span>'
        if seniority else ""
    )

    sal_html = (
        f'<span style="font-size:0.95rem;font-weight:700;color:#3ddc84;'
        f'font-family:\'JetBrains Mono\',monospace;white-space:nowrap;">{sal}</span>'
        if sal else ""
    )

    time_html = (
        f'<span style="font-size:0.75rem;color:#3a5040;white-space:nowrap;">{time_str}</span>'
        if time_str else ""
    )

    skills_html = ""
    if show_skills:
        raw = job.get("skills") or ""
        tags = [s.strip() for s in raw.split("|") if s.strip()][:7]
        if tags:
            skills_html = f'<div style="margin-top:0.5rem;">{"".join(_skill_tag(t) for t in tags)}</div>'

    return (
        f'<div style="background:#0f1310;border:1px solid #161e18;border-radius:10px;'
        f'padding:0.875rem 1.25rem;margin-bottom:0.45rem;">'
        f'<div style="display:flex;align-items:center;gap:0.875rem;">'
        f'{_logo(company, 42)}'
        f'<div style="flex:1;min-width:0;">'
        f'<div style="margin-bottom:0.2rem;">'
        f'<a href="{url}" target="_blank" style="font-size:0.94rem;font-weight:600;'
        f'color:#dce8e0;text-decoration:none;">{title}</a>{badge}'
        f'</div>'
        f'<div style="font-size:0.77rem;">'
        f'<span style="color:#6a9a7a;">{company}</span>'
        f'<span style="color:#283830;">&nbsp;&middot;&nbsp;</span>'
        f'<span style="color:#486054;">&#9737; {loc}</span>'
        f'<span style="color:#283830;">&nbsp;&middot;&nbsp;</span>'
        f'<span style="color:#486054;">&#9672; {cat}</span>'
        f'</div>'
        f'{skills_html}'
        f'</div>'
        f'<div style="display:flex;align-items:center;gap:1.1rem;flex-shrink:0;">'
        f'{sal_html}{time_html}'
        f'</div>'
        f'</div>'
        f'</div>'
    )

# ── Sidebar ───────────────────────────────────────────────────────────────────

with st.sidebar:

    # Logo
    st.markdown("""
<div style="padding:1.4rem 1.25rem 1.1rem;border-bottom:1px solid #192020;">
  <div style="display:flex;align-items:center;gap:0.7rem;margin-bottom:0.65rem;">
    <div style="width:40px;height:40px;border-radius:50%;background:#0c1e13;
                border:1.5px solid #2dbc74;display:flex;align-items:center;
                justify-content:center;flex-shrink:0;">
      <div style="width:14px;height:14px;border-radius:50%;border:1.5px solid #2dbc74;
                  opacity:0.7;display:flex;align-items:center;justify-content:center;">
        <div style="width:4px;height:4px;border-radius:50%;background:#3ddc84;"></div>
      </div>
    </div>
    <div>
      <div style="font-size:0.95rem;font-weight:700;color:#e0ebe4;letter-spacing:0.07em;
                  font-family:'JetBrains Mono',monospace;line-height:1.2;">LLMARKET</div>
      <div style="font-size:0.56rem;color:#2a4a34;letter-spacing:0.18em;
                  text-transform:uppercase;margin-top:2px;">SKILL RADAR</div>
    </div>
  </div>
  <div style="display:inline-flex;align-items:center;gap:0.3rem;
              background:#0c1e13;border:1px solid #1a3020;
              border-radius:20px;padding:2px 10px;">
    <div style="width:5px;height:5px;border-radius:50%;background:#3ddc84;"></div>
    <span style="font-size:0.6rem;color:#3ddc84;font-weight:600;letter-spacing:0.07em;">LIVE</span>
  </div>
</div>
""", unsafe_allow_html=True)

    # Nav header
    st.markdown("""
<div style="padding:0.9rem 1.25rem 0.35rem;">
  <div style="font-size:0.58rem;font-weight:700;color:#1e3028;
              letter-spacing:0.16em;text-transform:uppercase;">Navigation</div>
</div>
""", unsafe_allow_html=True)

    if "page" not in st.session_state:
        st.session_state.page = "Overview"

    _NAV = [
        ("Overview",       "▦"),
        ("Jobs",           "⊟"),
        ("Skills",         "◉"),
        ("Interview Prep", "◎"),
        ("Trends",         "◌"),
    ]
    for label, icon in _NAV:
        active = st.session_state.page == label
        st.markdown(f'<div class="{"nav-active " if active else ""}nav-wrap">', unsafe_allow_html=True)
        if st.button(f"{icon}  {label}", key=f"nav_{label}", use_container_width=True):
            st.session_state.page = label
            st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)

    # Divider
    st.markdown("<div style='height:1px;background:#192020;margin:0.5rem 0;'></div>",
                unsafe_allow_html=True)

    # Data Controls
    st.markdown("""
<div style="padding:0 1.25rem 0.4rem;">
  <div style="font-size:0.58rem;font-weight:700;color:#1e3028;
              letter-spacing:0.16em;text-transform:uppercase;">Data Controls</div>
</div>
""", unsafe_allow_html=True)

    all_locs = get_canonical_locations()
    selected_locs = st.multiselect(
        "Locations", options=all_locs, default=[],
        placeholder="All Locations", label_visibility="collapsed",
    )

    all_companies = sorted({e["name"] for e in COMPANY_REGISTRY})
    selected_companies = st.multiselect(
        "Companies", options=all_companies, default=[],
        placeholder="All Companies", label_visibility="collapsed",
    )

    all_roles = [
        "LLM / Generative AI", "Clinical AI", "ML Engineer", "Applied Scientist",
        "Data Scientist", "AI Platform / Infra", "Computer Vision",
        "NLP Engineer", "Data Engineer", "AI / ML",
    ]
    selected_roles = st.multiselect(
        "Job Categories", options=all_roles, default=[],
        placeholder="All Job Categories", label_visibility="collapsed",
    )

    min_sal_k = st.slider("Min Salary", 0, 400, 0, 10, format="$%dk")
    min_sal = min_sal_k * 1_000

    skill_search = st.text_input("Skill", placeholder="🔍  Filter by skill…",
                                 label_visibility="collapsed")

    # Divider
    st.markdown("<div style='height:1px;background:#192020;margin:0.5rem 0;'></div>",
                unsafe_allow_html=True)

    # Last updated
    all_jobs_meta = get_all_jobs()
    if all_jobs_meta:
        latest = max((j.get("fetched_at") or "" for j in all_jobs_meta), default="")
        if latest:
            try:
                dt_fmt = datetime.fromisoformat(latest[:19]).strftime("%b %d, %Y %I:%M %p")
            except Exception:
                dt_fmt = latest[:16]
            st.markdown(
                f'<div style="padding:0 1.25rem 0.5rem;font-size:0.67rem;'
                f'color:#2a4030;line-height:1.7;">Last updated<br>'
                f'<span style="color:#3a5848;">{dt_fmt}</span></div>',
                unsafe_allow_html=True,
            )

    st.markdown('<div class="refresh-btn" style="padding:0 0.75rem;">', unsafe_allow_html=True)
    fetch_clicked = st.button("↺  Refresh data", use_container_width=True, key="fetch_btn")
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown("""
<div style="padding:0.6rem 1.25rem 1.5rem;font-size:0.61rem;color:#192820;line-height:1.9;">
  Data from various sources<br>Results may vary
</div>
""", unsafe_allow_html=True)

# ── Fetch handler ─────────────────────────────────────────────────────────────

if fetch_clicked:
    with st.spinner("Fetching live job data…"):
        results: list[str] = []
        _area = st.empty()

        def _cb(name: str, stored: int, _skip: int) -> None:
            if stored:
                results.append(f"✓ {name} +{stored}")
            _area.markdown(
                '<div style="font-size:0.8rem;color:#3ddc84;font-family:monospace;">'
                + "<br>".join(results[-6:]) + "</div>",
                unsafe_allow_html=True,
            )

        summary = run_all_fetchers(progress_callback=_cb)
        _area.empty()
    st.toast(f"✓ {summary['total_stored']} jobs ingested", icon="✅")
    st.rerun()

# ── Filters ───────────────────────────────────────────────────────────────────

_filters: dict = {}
if min_sal > 0:           _filters["min_salary"] = min_sal
if selected_locs:         _filters["locations"] = selected_locs
if selected_companies:    _filters["companies"] = selected_companies
if selected_roles:        _filters["role_categories"] = selected_roles
if skill_search.strip():  _filters["skill"] = skill_search.strip()

# ── Shared data ───────────────────────────────────────────────────────────────

stats    = skills_summary_stats(_filters)
jobs_all = get_all_jobs(_filters)

# ── Auto-download Excel on first load of each session ────────────────────────
if not st.session_state.get("excel_auto_downloaded") and jobs_all:
    excel_bytes = _to_excel(jobs_all)
    b64 = base64.b64encode(excel_bytes).decode()
    fname = f"llmarket_{datetime.now().strftime('%Y%m%d')}.xlsx"
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    components.html(
        f'<a id="dl" href="data:{mime};base64,{b64}" download="{fname}"></a>'
        f'<script>document.getElementById("dl").click();</script>',
        height=0,
    )
    st.session_state.excel_auto_downloaded = True

# ── Plotly theme ──────────────────────────────────────────────────────────────

_CHART = dict(
    paper_bgcolor="rgba(0,0,0,0)",
    plot_bgcolor="#0f1310",
    font=dict(color="#6a8a74", family="Inter"),
    margin=dict(l=8, r=8, t=36, b=8),
    xaxis=dict(gridcolor="#161e18", zerolinecolor="#161e18"),
    yaxis=dict(gridcolor="#161e18", zerolinecolor="#161e18"),
)

# ── Stats bar (always visible) ────────────────────────────────────────────────

c1, c2, c3, c4 = st.columns(4)
c1.metric("TOTAL JOBS",    f"{stats['total_jobs']:,}",    delta="● Live openings")
c2.metric("UNIQUE SKILLS", f"{stats['unique_skills']:,}", delta="In demand")
with c3:
    st.markdown('<div class="salary-metric">', unsafe_allow_html=True)
    st.metric("AVG. PAY SALARY",
              f"${stats['avg_salary']:,.0f}" if stats["avg_salary"] else "—",
              delta="USD")
    st.markdown('</div>', unsafe_allow_html=True)
c4.metric("COMPANIES", f"{stats['companies']:,}", delta="Hiring now")

st.markdown("<div style='height:0.2rem'></div>", unsafe_allow_html=True)

# =============================================================================
# PAGE: OVERVIEW
# =============================================================================

if st.session_state.page == "Overview":

    # Hero
    st.markdown("""
<div style="text-align:center;padding:1.75rem 0 1rem;">
  <h1 style="font-size:2.4rem;font-weight:700;color:#dce8e0;line-height:1.2;margin-bottom:0.45rem;">
    Find your next&nbsp;<span style="color:#3ddc84;">high impact</span>&nbsp;role
  </h1>
  <p style="font-size:0.88rem;color:#344c3c;margin:0;">
    Real-time job intelligence for data, AI, and engineering talent.
  </p>
</div>
""", unsafe_allow_html=True)

    # Search bar
    _, sc, _ = st.columns([1, 5, 1])
    with sc:
        hero_q = st.text_input(
            "Search", placeholder="🔍  Search jobs, skills, companies…",
            key="hero_q", label_visibility="collapsed",
        )
    if hero_q.strip():
        _filters["skill"] = hero_q.strip()
        jobs_all = get_all_jobs(_filters)

    # Popular tags
    st.markdown("""
<div style="text-align:center;margin:0.6rem 0 1.1rem;">
  <span style="font-size:0.73rem;color:#243828;margin-right:0.5rem;">Popular:</span>
  <span style="display:inline-block;padding:3px 13px;background:#101714;border:1px solid #1e2a22;border-radius:20px;font-size:0.73rem;color:#547060;margin:2px 3px;">Data Engineer</span>
  <span style="display:inline-block;padding:3px 13px;background:#101714;border:1px solid #1e2a22;border-radius:20px;font-size:0.73rem;color:#547060;margin:2px 3px;">Machine Learning</span>
  <span style="display:inline-block;padding:3px 13px;background:#101714;border:1px solid #1e2a22;border-radius:20px;font-size:0.73rem;color:#547060;margin:2px 3px;">Remote</span>
  <span style="display:inline-block;padding:3px 13px;background:#101714;border:1px solid #1e2a22;border-radius:20px;font-size:0.73rem;color:#547060;margin:2px 3px;">AI Engineer</span>
  <span style="display:inline-block;padding:3px 13px;background:#101714;border:1px solid #1e2a22;border-radius:20px;font-size:0.73rem;color:#547060;margin:2px 3px;">LLM</span>
</div>
""", unsafe_allow_html=True)

    # Feed header + export
    cl, cm, cr = st.columns([2, 2, 1])
    with cl:
        st.markdown(
            '<div style="font-size:0.85rem;font-weight:600;color:#dce8e0;margin-bottom:0.6rem;">'
            '<span style="color:#3ddc84;">●</span>&nbsp; Live Job Feed</div>',
            unsafe_allow_html=True,
        )
    with cm:
        if jobs_all:
            st.markdown(
                f'<div style="font-size:0.73rem;color:#2a4030;padding-top:0.1rem;">'
                f'&#8595; Updated just now &nbsp;&middot;&nbsp; {len(jobs_all)} openings</div>',
                unsafe_allow_html=True,
            )
    with cr:
        if jobs_all:
            st.download_button(
                "⬇ Excel",
                _to_excel(jobs_all),
                file_name="llmarket_jobs.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

    if not jobs_all:
        st.markdown("""
<div style="text-align:center;padding:4rem 0;color:#2a4030;">
  <div style="font-size:2rem;margin-bottom:0.5rem;">◌</div>
  <div>No jobs yet. Click <b style="color:#3ddc84;">↺ Refresh data</b> in the sidebar.</div>
</div>
""", unsafe_allow_html=True)
    else:
        for job in jobs_all[:50]:
            st.markdown(_job_card(job, show_skills=True), unsafe_allow_html=True)
        if len(jobs_all) > 50:
            st.markdown(
                f'<div style="text-align:center;padding:1rem;font-size:0.78rem;color:#2a4030;">'
                f'Showing 50 of {len(jobs_all)} — use sidebar filters to narrow.</div>',
                unsafe_allow_html=True,
            )

# =============================================================================
# PAGE: JOBS
# =============================================================================

elif st.session_state.page == "Jobs":

    st.markdown(
        '<h2 style="font-size:1.25rem;font-weight:700;color:#dce8e0;margin-bottom:1.1rem;">'
        '⊟&nbsp; All Jobs</h2>',
        unsafe_allow_html=True,
    )

    if not jobs_all:
        st.info("No jobs match current filters.")
    else:
        ch, ce = st.columns([3, 1])
        with ch:
            st.markdown(
                f'<div style="font-size:0.78rem;color:#4a6050;margin-bottom:0.9rem;">'
                f'{len(jobs_all)} jobs</div>',
                unsafe_allow_html=True,
            )
        with ce:
            st.download_button(
                "⬇ Export Excel",
                _to_excel(jobs_all),
                file_name="llmarket_jobs.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

        for job in jobs_all:
            st.markdown(_job_card(job, show_skills=True), unsafe_allow_html=True)

# =============================================================================
# PAGE: SKILLS
# =============================================================================

elif st.session_state.page == "Skills":

    st.markdown(
        '<h2 style="font-size:1.25rem;font-weight:700;color:#dce8e0;margin-bottom:1.1rem;">'
        '◉&nbsp; Skill Intelligence</h2>',
        unsafe_allow_html=True,
    )

    freq_df = skill_frequency(_filters, top_n=40)

    if freq_df.empty:
        st.info("No skill data. Fetch jobs first.")
    else:
        tab_freq, tab_pay, tab_cat = st.tabs(["Frequency", "Top Paying", "By Category"])

        with tab_freq:
            col_chart, col_tbl = st.columns([3, 1])
            with col_chart:
                fig = px.bar(
                    freq_df.head(25), x="job_count", y="skill",
                    orientation="h", color="category",
                    color_discrete_sequence=[
                        "#3ddc84","#00b4d8","#f72585","#7b2d8b",
                        "#ff6b35","#ffd60a","#06d6a0","#8338ec",
                    ],
                    title="Most In-Demand Skills",
                    labels={"job_count": "Job Postings", "skill": ""},
                    height=600,
                )
                fig.update_layout(**_CHART, yaxis_autorange="reversed",
                                  title_font=dict(color="#e8f0eb", size=14))
                st.plotly_chart(fig, use_container_width=True)
            with col_tbl:
                st.dataframe(
                    freq_df[["skill", "category", "job_count", "pct"]].rename(
                        columns={"job_count": "Jobs", "pct": "%"}
                    ),
                    use_container_width=True, height=580,
                )
                st.download_button(
                    "⬇ Skills CSV",
                    freq_df.to_csv(index=False).encode(),
                    "skills.csv", "text/csv",
                    use_container_width=True,
                )

        with tab_pay:
            pay_df = top_paying_skills(_filters, top_n=20)
            if pay_df.empty:
                st.info("Need salary data across at least 2 jobs per skill.")
            else:
                fig2 = px.bar(
                    pay_df, x="median_salary", y="skill",
                    orientation="h", color="category",
                    color_discrete_sequence=["#3ddc84","#00b4d8","#f72585","#7b2d8b"],
                    title="Highest-Paying Skills (Median)",
                    labels={"median_salary": "Median Salary ($)", "skill": ""},
                    height=520, text="median_salary",
                )
                fig2.update_traces(
                    texttemplate="$%{text:,.0f}", textposition="outside",
                    textfont=dict(color="#3ddc84", size=10),
                )
                fig2.update_layout(**_CHART, yaxis_autorange="reversed",
                                   title_font=dict(color="#e8f0eb", size=14))
                st.plotly_chart(fig2, use_container_width=True)

        with tab_cat:
            fig3 = px.treemap(
                freq_df, path=["category", "skill"], values="job_count",
                color="job_count", color_continuous_scale="Greens",
                title="Skills by Category",
            )
            fig3.update_layout(**_CHART, title_font=dict(color="#e8f0eb", size=14))
            st.plotly_chart(fig3, use_container_width=True)

# =============================================================================
# PAGE: INTERVIEW PREP
# =============================================================================

elif st.session_state.page == "Interview Prep":

    st.markdown(
        '<h2 style="font-size:1.25rem;font-weight:700;color:#dce8e0;margin-bottom:0.3rem;">'
        '◎&nbsp; Interview Prep Intelligence</h2>'
        '<p style="font-size:0.78rem;color:#344c3c;margin-bottom:1.1rem;">'
        'Generated from real job postings · no AI inference · purely data-driven</p>',
        unsafe_allow_html=True,
    )

    all_role_opts = [
        None, "LLM / Generative AI", "Clinical AI", "ML Engineer",
        "Applied Scientist", "Data Scientist", "AI Platform / Infra",
        "Computer Vision", "NLP Engineer", "Data Engineer", "AI / ML",
    ]
    selected_role_prep = st.selectbox(
        "Focus on role", options=all_role_opts,
        format_func=lambda x: "All Roles" if x is None else x,
    )

    roadmap   = study_roadmap(role_category=selected_role_prep, top_n_skills=50)
    freq_prep = skill_frequency(_filters, top_n=60)

    if not roadmap:
        st.info("No data yet — fetch jobs first.")
    else:
        st.markdown("---")
        _ICONS = {
            "LLM Stack": "🤖", "ML/DL": "🧠", "Healthcare": "🏥",
            "MLOps": "⚙️", "Data Engineering": "🔧", "Vector Databases": "🗄️",
            "Inference": "⚡", "Cloud": "☁️", "Programming": "💻", "Governance": "🛡️",
        }
        cols = st.columns(2)
        for i, (cat, skills) in enumerate(roadmap.items()):
            with cols[i % 2]:
                st.markdown(
                    f'<div style="background:#0f1310;border:1px solid #161e18;'
                    f'border-radius:10px;padding:1rem;margin-bottom:0.7rem;">'
                    f'<div style="font-size:0.78rem;font-weight:600;color:#7aaa8a;'
                    f'margin-bottom:0.5rem;">{_ICONS.get(cat,"◆")} {cat}</div>'
                    + "".join(_skill_tag(s) for s in skills) + '</div>',
                    unsafe_allow_html=True,
                )

        st.markdown("---")
        st.markdown(
            '<div style="font-size:0.85rem;font-weight:600;color:#dce8e0;'
            'margin-bottom:0.6rem;">Most Common Skill Combinations</div>',
            unsafe_allow_html=True,
        )
        pairs_df = skill_combinations(_filters, top_n=15)
        if not pairs_df.empty:
            fig_p = px.bar(
                pairs_df, x="co_count",
                y=pairs_df["skill_a"] + " + " + pairs_df["skill_b"],
                orientation="h", height=420,
                color="co_count", color_continuous_scale="Greens",
                labels={"co_count": "Jobs Together", "y": ""},
            )
            fig_p.update_layout(**_CHART, yaxis_autorange="reversed",
                                coloraxis_showscale=False)
            st.plotly_chart(fig_p, use_container_width=True)

        if not freq_prep.empty:
            top_stack = (
                freq_prep[freq_prep["category"] == "LLM Stack"]["skill"].head(4).tolist()
                + freq_prep[freq_prep["category"] == "Healthcare"]["skill"].head(3).tolist()
                + freq_prep[freq_prep["category"] == "ML/DL"]["skill"].head(3).tolist()
                + freq_prep[freq_prep["category"] == "MLOps"]["skill"].head(3).tolist()
            )
            if top_stack:
                st.markdown(
                    '<div style="background:#0c1a12;border:1px solid #1a3022;'
                    'border-radius:10px;padding:1rem 1.25rem;margin-top:0.5rem;">'
                    '<div style="font-size:0.76rem;font-weight:600;color:#3ddc84;'
                    'margin-bottom:0.5rem;">⬡ Most common stack across healthcare AI jobs</div>'
                    + "".join(_skill_tag(s) for s in top_stack) + '</div>',
                    unsafe_allow_html=True,
                )

# =============================================================================
# PAGE: TRENDS
# =============================================================================

elif st.session_state.page == "Trends":

    st.markdown(
        '<h2 style="font-size:1.25rem;font-weight:700;color:#dce8e0;margin-bottom:0.3rem;">'
        '◌&nbsp; Trend Tracking</h2>'
        '<p style="font-size:0.78rem;color:#344c3c;margin-bottom:1.1rem;">'
        'Skill demand changes over time · snapshots taken on each fetch</p>',
        unsafe_allow_html=True,
    )

    days_sel = st.radio(
        "Window", [30, 60, 90],
        format_func=lambda d: f"{d} days", horizontal=True,
    )
    trend_df = skill_trend_df(days=days_sel)

    if trend_df.empty or trend_df["snapshot_date"].nunique() < 2:
        st.markdown("""
<div style="background:#0f1310;border:1px solid #161e18;border-radius:10px;
            padding:2rem;text-align:center;color:#4a6050;">
  <div style="font-size:1.5rem;margin-bottom:0.5rem;">◌</div>
  Trend data accumulates over time. Run
  <b style="color:#3ddc84;">↺ Refresh data</b> daily.
</div>
""", unsafe_allow_html=True)
        st.markdown("---")
        st.markdown(
            '<div style="font-size:0.85rem;font-weight:600;color:#dce8e0;'
            'margin-bottom:0.6rem;">Current Snapshot (Baseline)</div>',
            unsafe_allow_html=True,
        )
        freq_base = skill_frequency(top_n=25)
        if not freq_base.empty:
            fig_b = px.bar(
                freq_base, x="skill", y="job_count", color="category",
                color_discrete_sequence=[
                    "#3ddc84","#00b4d8","#f72585","#7b2d8b",
                    "#ff6b35","#ffd60a","#06d6a0","#8338ec",
                ],
                labels={"job_count": "Jobs", "skill": ""},
            )
            fig_b.update_layout(**_CHART, xaxis_tickangle=-45)
            st.plotly_chart(fig_b, use_container_width=True)
    else:
        growing_df = fastest_growing_skills(days=days_sel, top_n=15)
        if not growing_df.empty:
            st.markdown(
                '<div style="font-size:0.85rem;font-weight:600;color:#dce8e0;'
                'margin-bottom:0.6rem;">🚀 Fastest Growing Skills</div>',
                unsafe_allow_html=True,
            )
            fig_g = px.bar(
                growing_df, x="growth", y="skill", orientation="h",
                color="category",
                color_discrete_sequence=["#3ddc84","#00b4d8","#f72585","#7b2d8b"],
                height=420, text="growth",
                labels={"growth": "Δ Jobs", "skill": ""},
            )
            fig_g.update_traces(textposition="outside", textfont=dict(color="#3ddc84"))
            fig_g.update_layout(**_CHART, yaxis_autorange="reversed")
            st.plotly_chart(fig_g, use_container_width=True)

        st.markdown("---")
        top10 = (
            trend_df.groupby("canonical_name")["job_count"]
            .max().sort_values(ascending=False).head(10).index.tolist()
        )
        trend_top = trend_df[trend_df["canonical_name"].isin(top10)]
        if not trend_top.empty:
            st.markdown(
                '<div style="font-size:0.85rem;font-weight:600;color:#dce8e0;'
                'margin-bottom:0.6rem;">Top 10 Skills Over Time</div>',
                unsafe_allow_html=True,
            )
            fig_l = px.line(
                trend_top, x="snapshot_date", y="job_count",
                color="canonical_name", markers=True,
                labels={
                    "snapshot_date": "Date", "job_count": "Jobs",
                    "canonical_name": "Skill",
                },
            )
            fig_l.update_layout(**_CHART)
            st.plotly_chart(fig_l, use_container_width=True)

        with st.expander("Raw Snapshot Table"):
            pivot = trend_df.pivot_table(
                index="canonical_name", columns="snapshot_date",
                values="job_count", fill_value=0,
            )
            st.dataframe(pivot, use_container_width=True)
            st.download_button(
                "⬇ Export Trends CSV",
                pivot.to_csv().encode(), "trends.csv", "text/csv",
            )
